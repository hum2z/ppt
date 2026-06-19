"""Offline tests for the deck engine (no API key required).

Builds a small reference deck with deliberate formatting, runs the
extract -> rewrite round-trip, and asserts that text is swapped while font
styling (size, bold, colour) is preserved.
"""

import os
import sys

from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.deck import apply_translations, extract_blueprint  # noqa: E402


def build_reference(path: str) -> None:
    prs = Presentation()

    # Slide 1: title layout
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "The Roman Empire"
    s1.placeholders[1].text = "A brief overview"

    # Slide 2: title + bullets, with custom formatting on the first bullet
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Key Facts"
    body = s2.placeholders[1].text_frame
    body.text = "Founded in 27 BC"
    run0 = body.paragraphs[0].runs[0]
    run0.font.size = Pt(28)
    run0.font.bold = True
    run0.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    p = body.add_paragraph()
    p.text = "Spanned three continents"
    p.level = 1

    # Slide 3: a table
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Comparison"
    table = s3.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    table.cell(0, 0).text = "Category"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Population"
    table.cell(1, 1).text = "Millions"

    prs.save(path)


def main() -> int:
    ref = "/tmp/reference.pptx"
    out = "/tmp/output.pptx"
    build_reference(ref)

    bp = extract_blueprint(ref)
    assert len(bp.slides) == 3, f"expected 3 slides, got {len(bp.slides)}"
    ids = [e.id for s in bp.slides for e in s.elements]
    assert len(ids) == len(set(ids)), "element ids must be unique"
    assert bp.element_count >= 8, f"too few elements: {bp.element_count}"

    roles = {e.id: e.role for s in bp.slides for e in s.elements}
    assert any(r == "title" for r in roles.values()), "expected at least one title role"
    assert any(r == "table-cell" for r in roles.values()), "expected table-cell roles"

    # Fabricate translations (what Claude would return) for every id.
    translations = {e.id: f"NEW::{e.text}" for s in bp.slides for e in s.elements}
    replaced = apply_translations(ref, out, translations)
    assert replaced == bp.element_count, f"replaced {replaced} of {bp.element_count}"

    # Re-extract output and confirm every text was swapped.
    bp2 = extract_blueprint(out)
    for s in bp2.slides:
        for e in s.elements:
            assert e.text.startswith("NEW::"), f"not replaced: {e.text!r}"

    # Confirm formatting survived on the styled bullet.
    prs = Presentation(out)
    styled = None
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    if "Founded in 27 BC" in para.text:
                        styled = para.runs[0]
    assert styled is not None, "styled bullet not found"
    assert styled.font.bold is True, "bold lost"
    assert styled.font.size == Pt(28), f"size lost: {styled.font.size}"
    assert styled.font.color.rgb == RGBColor(0xC0, 0x39, 0x2B), "colour lost"

    print("OK: extraction, replacement, and formatting preservation all pass.")
    print(f"  slides={len(bp.slides)} elements={bp.element_count} replaced={replaced}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
